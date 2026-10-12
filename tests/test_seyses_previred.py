import io
import unittest

import pandas as pd
from openpyxl import load_workbook

from seyses_previred import generar_excel_formato_previred


class PreviredExportTests(unittest.TestCase):
    def test_export_returns_valid_xlsx_with_expected_sheet_and_columns(self):
        source = pd.DataFrame([{
            "RUT": "12.345.678-5",
            "Nombre PDF": "Ana Pérez",
            "Renta AFP": 1000000,
            "Cotizacion AFP": 110000,
        }])
        content = generar_excel_formato_previred(source)
        self.assertTrue(content.startswith(b"PK"))
        workbook = load_workbook(io.BytesIO(content), read_only=True)
        self.assertIn("Previred_Report", workbook.sheetnames)
        sheet = workbook["Previred_Report"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        self.assertIn("RUT", headers)
        self.assertIn("Renta AFP", headers)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
